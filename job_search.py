"""Job search automation tool using Playwright, Gemini API, and Pandas."""

import datetime
import json
import os
import random
import re
import sys
import time
import urllib.parse

from dotenv import load_dotenv
from google import genai
from google.genai import types
from openpyxl.styles import Font, PatternFill
import pandas as pd
from playwright.sync_api import sync_playwright

load_dotenv()

# グローバル定数の定義
MAX_RETRIES = 3
NG_COMPANY_WORDS = [
    "社名非公開",
    "非公開",
    "匿名",
    "某大手企業",
    "(事業所の意向により公開していません)",
    "事業所の意向により公開していません",
]


def load_json_config(file_path="config.json"):
    """Loads external search criteria configuration file (JSON)."""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError as e:
        raise FileNotFoundError(
            f"Configuration file {file_path} not found. Execution stopped."
        ) from e
    except json.JSONDecodeError as e:
        raise ValueError(
            f"Configuration file {file_path} is not in a valid JSON format. Execution stopped."
        ) from e


def inspect_companies_batch_with_gemini(matched_jobs):
    """Batch inspects matched companies with Gemini API for SES/dispatch classification."""
    api_key = os.getenv("GEMINI_API_KEY")
    model_name = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")

    if not api_key:
        raise ValueError(
            "❌ Critical Error: GEMINI_API_KEY is not set in the environment variables."
        )

    company_entries = []
    for idx, job in enumerate(matched_jobs, 1):
        summary = job.get("business_summary", "")[:100]
        company_entries.append(
            f"ID:{idx} | 企業名: {job.get('company')} | 職種: {job.get('title')} | 概要: {summary}"
        )

    companies_text = "\n".join(company_entries)

    prompt = f"""
以下の企業リストについて、
許認可（労働者派遣事業や有料職業紹介事業）の記載が少しでも推測される場合、または求人内容が実質的にSESや客先常駐である場合は、迷わず true（派遣・SES）にすること。

【対象企業リスト】
{companies_text}

以下の構造化JSON配列形式でのみ出力してください（IDごとに返却）。
[
  {{
    "id": 1,
    "company_name": "企業名",
    "business_summary": "この企業の一般的な事業内容（100文字程度のわかりやすい日本語要約）",
    "is_ses_or_dispatch": true または false (労働者派遣事業、有料職業紹介事業、客先常駐、SES、エンジニア派遣、アウトソーシングをメインとしている場合は true),
    "reason": "判定の短い理由"
  }},
  ...
]
"""
    client = genai.Client(api_key=api_key)

    for attempt in range(1, MAX_RETRIES + 1):
        print(
            f"🚀 Sending {len(matched_jobs)} jobs to AI ({model_name})... (Attempt: {attempt}/{MAX_RETRIES})"
        )

        try:
            response = client.models.generate_content(
                model=model_name,
                contents=prompt,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json", temperature=0.2
                ),
            )
        except Exception as e:
            print(
                f"⚠️ Network connection or API call failed (Attempt {attempt}/{MAX_RETRIES}): {e}"
            )
            if attempt < MAX_RETRIES:
                time.sleep(attempt * 3)
                continue
            else:
                print("❌ AI connection failed. Proceeding without AI classification.")
                for job in matched_jobs:
                    job["is_ses_or_dispatch"] = False
                return matched_jobs

        res_text = response.text.strip()
        data_list = json.loads(res_text)
        result_map = {item.get("id"): item for item in data_list}

        for idx, job in enumerate(matched_jobs, 1):
            res = result_map.get(idx, {})
            is_ses = res.get("is_ses_or_dispatch", False)
            summary = res.get("business_summary", job.get("business_summary"))

            job["is_ses_or_dispatch"] = is_ses
            if is_ses:
                job["business_summary"] = f"【派遣・SES業】{summary}"
            else:
                job["business_summary"] = summary
        return matched_jobs

    return matched_jobs


def fetch_google_jobs(keyword_query, location):
    """Job scraping function using Playwright."""
    loc_str = (
        str(location).strip()
        if location and str(location).lower() != "null"
        else ""
    )
    print(
        f"\n🔍 Scrapes job postings from Google Jobs. [{keyword_query} × {loc_str or '全国'}]"
    )

    raw_query = f"{keyword_query} 求人 {loc_str}".strip()
    encoded_query = urllib.parse.quote(raw_query)
    google_search_url = (
        f"https://www.google.com/search?q={encoded_query}&hl=ja&gl=jp"
    )

    job_results = []
    screenshot_path = f"debug_jobs_{loc_str}.png"

    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(
                headless=False,
                args=[
                    "--disable-blink-features=AutomationControlled",
                    "--disable-infobars",
                    "--no-sandbox",
                    "--disable-setuid-sandbox",
                    "--disable-dev-shm-usage",
                    "--start-maximized",
                ],
            )

            context = browser.new_context(
                user_agent=(
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/122.0.0.0 Safari/537.36"
                ),
                viewport={"width": 1920, "height": 1080},
                locale="ja-JP",
                timezone_id="Asia/Tokyo",
                permissions=["geolocation"],
            )

            context.add_init_script("""
                Object.defineProperty(navigator, 'webdriver', {
                    get: () => undefined
                });
                window.navigator.chrome = {
                    runtime: {}
                };
                Object.defineProperty(navigator, 'plugins', {
                    get: () => [1, 2, 3, 4, 5]
                });
            """)

            page = context.new_page()

            print(f"📡 Now Scraping: {google_search_url}")
            page.goto(
                google_search_url,
                timeout=30000,
                wait_until="domcontentloaded",
            )
            page.wait_for_timeout(random.randint(2000, 3000))
            page.mouse.move(random.randint(100, 500), random.randint(100, 500))

            content = page.locator("body").inner_text()
            if (
                "ロボット" in content
                or "unusual traffic" in content
                or "CAPTCHA" in content
            ):
                print("🚨 Warning: Google bot block (CAPTCHA) detected.")
                print(
                    "💡 Please solve the puzzle or authentication manually on the browser screen."
                )
                input(
                    "👉 Please clear the robot verification on the browser, then press Enter..."
                )

            try:
                more_jobs_btn = page.query_selector(
                    "a[href*='udm=jobs'], div[role='button']:has-text('求人'), a:has-text('求人')"
                )
                if more_jobs_btn:
                    more_jobs_btn.click()
                    page.wait_for_timeout(3000)
            except Exception:
                pass

            page.screenshot(path=screenshot_path)
            body_text = page.locator("body").inner_text()
            browser.close()

        lines = [line.strip() for line in body_text.split("\n") if line.strip()]

        for i, line in enumerate(lines):
            if "提供元:" in line or "via" in line:
                title = (
                    lines[i - 2]
                    if i >= 2
                    else (lines[i - 1] if i >= 1 else "N/A")
                )
                company = lines[i - 1] if i >= 1 else "Confidential"
                source_raw = line

                loc = loc_str if loc_str else "Not specified"
                salary = "See description"

                desc_chunks = []
                for offset in range(1, 10):
                    if i + offset < len(lines):
                        sub_line = lines[i + offset]
                        desc_chunks.append(sub_line)
                        if any(
                            c in sub_line
                            for c in ["円", "万", "月", "年", "時給"]
                        ):
                            if re.search(r"\d", sub_line):
                                salary = sub_line
                        elif any(
                            pref in sub_line
                            for pref in [
                                "大阪",
                                "京都",
                                "兵庫",
                                "滋賀",
                                "奈良",
                                "和歌山",
                                "〒",
                            ]
                        ):
                            clean_loc = re.split(r"[•|·]|提供元:", sub_line)[
                                0
                            ].strip()
                            if clean_loc:
                                loc = clean_loc

                description_text = (
                    " / ".join(desc_chunks)
                    if desc_chunks
                    else f"{keyword_query}に関する事業・開発業務"
                )

                ng_words = [
                    "もっと見る",
                    "雇用形態",
                    "掲載日",
                    "地域を選択",
                    "保存した求人情報",
                    "フォロー中",
                    "求人",
                    "検索結果",
                ]
                if title and not any(ng in title for ng in ng_words):
                    job_results.append(
                        {
                            "title": title,
                            "company": company,
                            "business_summary": description_text,
                            "source": source_raw,
                            "location": loc,
                            "salary": salary,
                            "search_location": loc_str,
                        }
                    )

        if job_results:
            unique_jobs = []
            seen = set()
            for j in job_results:
                identifier = f"{j['title']}_{j['company']}"
                if identifier not in seen:
                    seen.add(identifier)
                    unique_jobs.append(j)

            print(
                f"✅ {loc_str or '全国'} エリアより {len(unique_jobs)} 件の求人を抽出しました。"
            )
            return unique_jobs

    except Exception as e:
        print(f"⚠️ [{loc_str}] Fetch exception occurred. ({e})。")

    print(f"⚠️ [{loc_str}] Real-time fetch failed. Loading backup data...")
    return [
        {
            "title": "エンジニア求人（開発業務）",
            "company": "株式会社ReaLight",
            "business_summary": (
                "自社Webサービスの企画・受託開発・プログラミングスクール運営業務"
            ),
            "source": "提供元: engage",
            "location": loc_str or "大阪府",
            "salary": "年 ¥420万〜¥520万",
        }
    ]


def clean_source_name(source_str):
    """Extracts the clean source site name from the raw source string."""
    if not source_str:
        return "Unknown"
    cleaned = source_str.replace("提供元:", "").replace("via", "").strip()
    if "•" in cleaned:
        parts = [p.strip() for p in cleaned.split("•") if p.strip()]
        for part in parts:
            if not any(
                pref in part
                for pref in [
                    "大阪",
                    "京都",
                    "兵庫",
                    "滋賀",
                    "奈良",
                    "和歌山",
                    "〒",
                ]
            ):
                cleaned = part
                break
    return cleaned


def parse_salary_to_monthly(salary_str):
    """Extracts minimum salary from raw string and converts to monthly value."""
    if not salary_str or "条件欄参照" in salary_str:
        return 0

    try:
        cleaned = (
            salary_str.replace(",", "")
            .replace("￥", "")
            .replace("¥", "")
            .replace("～", "~")
        )
        lower_part = cleaned.split("~")[0].split("-")[0].strip()

        match = re.search(r"([\d.]+)\s*(万|円)?", lower_part)
        if match:
            num = float(match.group(1))
            unit = match.group(2)

            val = num * 10_000 if unit == "万" else num

            if "年" in salary_str or "年俸" in salary_str:
                val = val / 12

            return int(val)
    except Exception:
        pass

    return 0


def filter_jobs_by_profile(
    job_list,
    keyword_config,
    exclude_private_company=True,
    target_sources=None,
    min_salary=0,
):
    """Primary filtering function for job data, handles keyword logic."""
    if isinstance(keyword_config, dict):
        kw_values = keyword_config.get("values", [])
        kw_mode = keyword_config.get("mode", "OR").upper()
    else:
        kw_values = (
            keyword_config
            if isinstance(keyword_config, list)
            else [str(keyword_config)]
        )
        kw_mode = "OR"

    target_keywords = [
        str(kw).strip().lower() for kw in kw_values if kw and str(kw).strip()
    ]

    filtered_jobs = []

    for job in job_list:
        company_name = job.get("company", "")
        source_name = clean_source_name(job.get("source", ""))

        if target_sources:
            matched_source = any(
                src.lower() in source_name.lower() for src in target_sources
            )
            if not matched_source:
                continue

        if min_salary and min_salary > 0:
            monthly_salary = parse_salary_to_monthly(job.get("salary", ""))
            if 0 < monthly_salary < min_salary:
                continue

        text_content = f"{job.get('title', '')} {company_name} {job.get('business_summary', '')}".lower()

        if exclude_private_company:
            if any(ng_word in company_name for ng_word in NG_COMPANY_WORDS):
                continue

        if target_keywords:
            if kw_mode == "AND":
                is_keyword_matched = all(
                    kw in text_content for kw in target_keywords
                )
                hits = target_keywords if is_keyword_matched else []
            else:
                hits = [kw for kw in target_keywords if kw in text_content]
                is_keyword_matched = len(hits) >= 1

            if not is_keyword_matched:
                continue
        else:
            hits = []

        job["matched_keywords"] = hits
        job["source"] = source_name
        filtered_jobs.append(job)

    return filtered_jobs


def export_to_excel(jobs_data, keyword_display, locations):
    """Exports structured job data into an Excel spreadsheet."""
    if not jobs_data:
        print("⚠️ No job data available to export.")
        return

    if isinstance(locations, list):
        loc_str = "_".join([str(l) for l in locations if l])
    else:
        loc_str = (
            str(locations).strip()
            if locations and str(locations).lower() != "null"
            else "全国"
        )

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"job_results_{keyword_display}_{loc_str}_{timestamp}.xlsx"

    records = []
    for row_job in jobs_data:
        records.append(
            {
                "検索地域": row_job.get("search_location", ""),
                "職種タイトル": row_job.get("title", ""),
                "企業名": row_job.get("company", ""),
                "事業内容 / 掲載概要": row_job.get("business_summary", ""),
                "勤務地詳細": row_job.get("location", ""),
                "給与条件": row_job.get("salary", ""),
                "検知キーワード": ", ".join(
                    row_job.get("matched_keywords", [])
                ),
                "掲載サイト": row_job.get("source", ""),
                "_is_ses": row_job.get("is_ses_or_dispatch", False),
            }
        )

    df = pd.DataFrame(records)

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        export_df = df.drop(columns=["_is_ses"])
        export_df.to_excel(writer, sheet_name="求人検索結果", index=False)
        worksheet = writer.sheets["求人検索結果"]

        worksheet.auto_filter.ref = worksheet.dimensions

        header_font = Font(name="Meiryo", size=11, bold=True, color="000000")
        header_fill = PatternFill(
            start_color="FCD5B4", end_color="FCD5B4", fill_type="solid"
        )

        for col_idx in range(1, len(export_df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill

        gray_fill = PatternFill(
            start_color="C0C0C0", end_color="C0C0C0", fill_type="solid"
        )
        for row_idx, is_ses in enumerate(df["_is_ses"], start=2):
            if is_ses:
                for col_idx in range(1, len(export_df.columns) + 1):
                    worksheet.cell(row=row_idx, column=col_idx).fill = gray_fill

        for col in worksheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            worksheet.column_dimensions[col_letter].width = max(
                min(max_len + 5, 50), 12
            )


def cleanup_debug_screenshots(locations):
    """Cleans up and removes temporary PNG screenshots saved during debugging."""
    loc_list = locations if isinstance(locations, list) else [locations]

    for loc in loc_list:
        loc_str = (
            str(loc).strip()
            if loc and str(loc).lower() != "null"
            else ""
        )
        screenshot_path = f"debug_jobs_{loc_str}.png"
        if os.path.exists(screenshot_path):
            try:
                os.remove(screenshot_path)
                print(f"  🗑️ Deleted successfully: {screenshot_path}")
            except Exception as e:
                print(f"  ⚠️ Failed to delete ({screenshot_path}): {e}")


def search_jobs_test():
    """Main routine to load configuration, fetch, filter, and export jobs."""
    json_config = load_json_config("config.json")
    conditions = json_config["search_conditions"]

    keyword_config = conditions.get("keyword", "")
    if isinstance(keyword_config, dict):
        kw_values = keyword_config.get("values", [])
        kw_mode = keyword_config.get("mode", "OR").upper()
    else:
        kw_values = (
            keyword_config
            if isinstance(keyword_config, list)
            else [str(keyword_config)]
        )
        kw_mode = "OR"

    if kw_values:
        keyword_query = " ".join([str(k) for k in kw_values if k])
        keyword_display = f"{kw_mode}_" + "_".join(
            [str(k) for k in kw_values if k]
        )
    else:
        keyword_query = ""
        keyword_display = "all"

    loc_config = conditions.get("location", [])
    if isinstance(loc_config, dict):
        location_values = loc_config.get("values", [])
        loc_mode = loc_config.get("mode", "OR").upper()
    else:
        location_values = (
            loc_config
            if isinstance(loc_config, list)
            else [str(loc_config)]
        )
        loc_mode = "OR"

    target_sources = conditions.get("target_sources", [])
    min_salary = conditions.get("min_salary", 0)

    if isinstance(location_values, str):
        location_list = [location_values]
    elif isinstance(location_values, list):
        location_list = location_values
    else:
        location_list = [None]

    print("⚙️ 検索条件を読み込みました:")
    print(f"   [キーワード]: {keyword_query} (モード: {kw_mode})")
    print(
        f"   [対象エリア]: {', '.join([str(l) for l in location_list if l])} (モード: {loc_mode})"
    )
    if target_sources:
        print(f"   [対象掲載先]: {', '.join(target_sources)}")
    if min_salary:
        print(f"   [最低給与条件]: 月給換算 {min_salary:,} 円以上")

    all_raw_jobs = []

    for loc in location_list:
        raw_jobs = fetch_google_jobs(keyword_query=keyword_query, location=loc)
        all_raw_jobs.extend(raw_jobs)

        if loc != location_list[-1]:
            sleep_sec = random.uniform(6.0, 10.0)
            print(
                f"⏳ 連続アクセスのブロックを防ぐため、次のエリアまで {sleep_sec:.1f} 秒待機します..."
            )
            time.sleep(sleep_sec)

    unique_raw_jobs = []
    seen = set()
    for j in all_raw_jobs:
        identifier = f"{j['title']}_{j['company']}"
        if identifier not in seen:
            seen.add(identifier)
            unique_raw_jobs.append(j)

    matched_jobs = filter_jobs_by_profile(
        unique_raw_jobs,
        keyword_config=keyword_config,
        exclude_private_company=True,
        target_sources=target_sources,
        min_salary=min_salary,
    )

    print(
        "\n==============================================================================="
    )
    print(
        f" 📊 一次選別完了: 取得 {len(unique_raw_jobs)} 件 ──[社名・掲載先・給与・KWフィルタ]──> 通過: {len(matched_jobs)} 件"
    )
    print(
        "==============================================================================="
    )

    if not matched_jobs:
        print("⚠️ No matching public job listings found.")
        cleanup_debug_screenshots(location_list)
        return

    verified_jobs = inspect_companies_batch_with_gemini(matched_jobs)

    export_to_excel(
        verified_jobs,
        keyword_display=keyword_display,
        locations=location_list,
    )

    cleanup_debug_screenshots(location_list)


if __name__ == "__main__":
    search_jobs_test()